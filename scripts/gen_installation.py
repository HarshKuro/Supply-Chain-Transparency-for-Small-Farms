import os
import sys
from docx.shared import Inches, Pt, RGBColor
from doc_utils import (
    init_document, add_title_header, add_heading_1, add_heading_2, add_heading_3,
    add_body_paragraph, add_bullet_point, add_callout, add_diagram_box,
    add_styled_table, add_screenshot
)

def build_installation_doc():
    doc = init_document("Technical Installation & Deployment Manual", "Technical Manual")
    add_title_header(
        doc,
        title="Technical Installation & Deployment Manual",
        subtitle="Complete engineering guide for Firebase project creation, Firestore security rules deployment, Australian database seeding, and production hosting.",
        category_tag="System Engineering Guide"
    )

    # 1. System Architecture & Prerequisites
    add_heading_1(doc, "1. Architecture Overview & Prerequisites")
    add_body_paragraph(
        doc,
        "The AgriTrace application is engineered as a lightweight, high-performance serverless web platform utilizing modern ES Modules (ECMAScript 6+), Vanilla CSS3 variables, Google Firebase Authentication, and Cloud Firestore. It requires zero compilation, bundlers, or heavy runtime frameworks, ensuring rapid deployment and extreme reliability."
    )

    prereq_headers = ["Prerequisite Component", "Minimum Required Version", "Purpose in Architecture"]
    prereq_rows = [
        ["Web Browser", "Chrome 110+, Edge 110+, Safari 16+", "Client-side ES Modules & Lucide iconography"],
        ["Local HTTP Server", "Python 3.8+ OR Node.js v18+", "Serving ES Modules without CORS local origin errors"],
        ["Google Firebase Account", "Active Google Cloud / Firebase console", "Authentication and real-time NoSQL Firestore database"],
        ["Operating System", "Windows 10/11, macOS 12+, Linux Ubuntu 20+", "Cross-platform shell compatibility"]
    ]
    add_styled_table(doc, prereq_headers, prereq_rows, [1.8, 2.0, 2.7])

    add_screenshot(doc, "13_browser_installation_manual.png", "AgriTrace Technical Installation & Deployment Manual Interface (installation-manual.html)")

    # 2. Step-by-Step Installation Procedures
    add_heading_1(doc, "2. Step-by-Step Installation & Firebase Setup")
    
    add_heading_2(doc, "Step 2.1: Firebase Project Creation")
    add_bullet_point(doc, "Open the Firebase Console at https://console.firebase.google.com and click 'Add project'.", "1. ")
    add_bullet_point(doc, "Enter Project Name: AgriTrace-Australia (or your organization's preference).", "2. ")
    add_bullet_point(doc, "Enable Google Analytics (optional) and click 'Create Project'.", "3. ")

    add_heading_2(doc, "Step 2.2: Enable Firebase Authentication")
    add_bullet_point(doc, "Navigate to Build > Authentication in the left navigation sidebar.", "1. ")
    add_bullet_point(doc, "Click 'Get Started', then select 'Email/Password' under Native Providers.", "2. ")
    add_bullet_point(doc, "Toggle 'Enable' on for Email/Password and click 'Save'.", "3. ")

    add_heading_2(doc, "Step 2.3: Initialize Cloud Firestore")
    add_bullet_point(doc, "Navigate to Build > Firestore Database and click 'Create database'.", "1. ")
    add_bullet_point(doc, "Choose location: australia-southeast1 (Sydney) for lowest regional latency.", "2. ")
    add_bullet_point(doc, "Select 'Start in production mode' to enforce security rules.", "3. ")

    add_heading_2(doc, "Step 2.4: Deploy Cloud Firestore Security Rules")
    add_body_paragraph(
        doc,
        "Copy and paste the project's firestore.rules into the Firestore Rules tab in the console. These rules enforce role validation and wholesaler accreditation checks:"
    )

    rules_lines = [
        "rules_version = '2';",
        "service cloud.firestore {",
        "  match /databases/{database}/documents {",
        "    function isSignedIn() { return request.auth != null; }",
        "    function getUserData() { return get(/databases/$(database)/documents/users/$(request.auth.uid)).data; }",
        "    function hasRole(r) { return isSignedIn() && getUserData().role == r; }",
        "    function isApprovedWholesaler() { return hasRole('wholesaler') && getUserData().approved == true; }",
        "    function isAdmin() { return hasRole('admin'); }",
        "    ",
        "    match /users/{userId} {",
        "      allow read: if isSignedIn();",
        "      allow create: if isSignedIn() && request.auth.uid == userId;",
        "      allow update: if isSignedIn() && (request.auth.uid == userId || isAdmin());",
        "    }",
        "    match /products/{productId} {",
        "      allow read: if true;",
        "      allow write: if hasRole('farmer') || isAdmin();",
        "    }",
        "    match /orders/{orderId} {",
        "      allow read: if isSignedIn();",
        "      allow create: if hasRole('customer');",
        "      allow update: if hasRole('farmer') || isApprovedWholesaler() || hasRole('driver') || isAdmin();",
        "    }",
        "  }",
        "}"
    ]
    add_diagram_box(doc, "Firestore Role Security Schema", rules_lines)

    add_heading_2(doc, "Step 2.5: Link Firebase Credentials")
    add_body_paragraph(
        doc,
        "In Firebase Project Settings > General > Your apps, register a Web App ('AgriTrace Web') and copy the firebaseConfig object into js/firebase-config.js in your workspace root."
    )

    # 3. Database Seeding & Australian Regional Entities
    add_heading_1(doc, "3. Executing the Australian Database Seeder (seed.html)")
    add_body_paragraph(
        doc,
        "To populate authentic Australian regional entities, certified harvests with AUD pricing, transit dispatches, and pre-configured accounts:"
    )
    add_bullet_point(doc, "Launch the local web server and open http://localhost:8000/seed.html in your browser.", "1. ")
    add_bullet_point(doc, "Click 'Run Database Seeder'. The script creates 6 authenticated stakeholders, 7 regional harvests (Yarra Valley Strawberries, Avocados, Olive Oil), sample orders, and customer ratings.", "2. ")
    add_bullet_point(doc, "Watch the real-time terminal log window confirm each write with green checkmarks.", "3. ")

    # 4. Running Locally
    add_heading_1(doc, "4. Launching the Local Web Server")
    add_body_paragraph(
        doc,
        "Because AgriTrace utilizes native ES Module imports (import ... from './module.js'), modern browsers require pages to be served over HTTP/HTTPS rather than the file:// protocol."
    )
    add_bullet_point(doc, "Python HTTP Server: python -m http.server 8000", "Option A: ")
    add_bullet_point(doc, "Node.js HTTP Server: npx http-server -p 8000", "Option B: ")
    add_bullet_point(doc, "Windows Automated Batch: Double-click run.bat in the root folder", "Option C: ")
    add_body_paragraph(
        doc,
        "Open your web browser and navigate to http://localhost:8000/. The AgriTrace home page will load immediately."
    )

    # 5. Automated Verification & Testing
    add_heading_1(doc, "5. Automated Verification & System Audit Suite")
    add_body_paragraph(
        doc,
        "The repository contains an automated Node.js test suite verifying 31 critical checkpoints across unit logic, system workflows, and production SEO compliance:"
    )

    test_lines = [
        "> node tests/unit_tests.js",
        "  ✔ 8/8 Unit Tests Passed (Form validation, quantities, totals, AUD currency formatting)",
        "",
        "> node tests/system_integration_tests.js",
        "  ✔ 14/14 System Tests Passed (Auth sync, stock deduction, lifecycle, UAT roles)",
        "",
        "> node tests/seo_audit.js",
        "  ✔ 9/9 SEO Audits Passed (robots.txt, sitemap.xml, Open Graph tags, single H1)",
        "",
        "SUMMARY: 31/31 TESTS PASSED (100% SUCCESS RATE)"
    ]
    add_diagram_box(doc, "Automated Test Suite Execution Log", test_lines)

    add_screenshot(doc, "01_terminal_unit_tests.png", "Terminal Execution: Unit Test Suite (8/8 Passed)")
    add_screenshot(doc, "03_terminal_system_integration_tests.png", "Terminal Execution: System Integration & UAT Test Suite (14/14 Passed)")
    add_screenshot(doc, "06_browser_automated_test_runner.png", "Interactive In-Browser Test Suite Runner (test_runner.html)")

    output_path = os.path.join("docx", "installationmanual.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_installation_doc()
