import os
import sys
from docx.shared import Inches, Pt, RGBColor
from doc_utils import (
    init_document, add_title_header, add_heading_1, add_heading_2, add_heading_3,
    add_body_paragraph, add_bullet_point, add_callout,
    add_styled_table, add_screenshot
)

def build_wholesaler_approval_doc():
    doc = init_document("Wholesaler Accreditation & Approval Policy", "Regulatory Policy")
    add_title_header(
        doc,
        title="Wholesaler Accreditation & Administrative Approval Workflow",
        subtitle="A regulatory and technical specification detailing why and how produce wholesalers are vetted and approved by administrators before receiving audit authority.",
        category_tag="Regulatory & Compliance Policy"
    )

    # 1. Background & Regulatory Imperative
    add_heading_1(doc, "1. Executive Summary & Regulatory Imperative")
    add_body_paragraph(
        doc,
        "In traditional agricultural supply chains, wholesalers exert substantial market power. However, granting unverified digital accounts unrestricted authority to audit and verify perishable food shipments exposes small family farms to severe fraud risks, counterfeit certification, and non-compliance with the Australia New Zealand Food Standards Code (FSANZ) and the Australian Biosecurity Act 2015."
    )
    add_body_paragraph(
        doc,
        "To safeguard small growers and maintain the integrity of AgriTrace's transparent custody chain, wholesalers CANNOT automatically inspect or certify stock upon registration. Instead, every wholesaler must undergo mandatory Administrative Accreditation."
    )

    add_callout(
        doc,
        "LEGAL & SAFETY PRINCIPLE: Only accredited wholesalers who have verified Australian business identities, compliant cold storage facilities, and food safety training may certify farm batches into commercial transit.",
        "Regulatory Biosecurity Requirement",
        "warning"
    )

    # 2. End-to-End Approval Workflow (REAL IMAGE)
    add_heading_1(doc, "2. End-to-End Accreditation Lifecycle Flow Diagram")
    add_body_paragraph(
        doc,
        "Below is the complete visual workflow chart governing wholesaler account onboarding, permission gating, administrative review, and status revocation:"
    )

    # Embedded high-res visual image diagram (NO text signs)
    add_screenshot(doc, "flow_wholesaler_accreditation.png", "Wholesaler Accreditation & Administrative Approval Governance Lifecycle", width_inches=6.2)

    add_screenshot(doc, "15_browser_admin_approvals.png", "Administrator Wholesaler Accreditation & Approvals Hub (admin/approvals.html)", width_inches=6.0)

    # 3. Functional Permissions Matrix
    add_heading_1(doc, "3. Functional Permissions Matrix: Pending vs. Approved")
    add_body_paragraph(
        doc,
        "The table below contrasts the exact system capabilities of unaccredited versus accredited wholesaler accounts:"
    )

    perm_headers = ["Portal Feature / Action", "Pending Wholesaler", "Approved Wholesaler", "Enforcement Layer"]
    perm_rows = [
        ["Log into Wholesaler Portal", "ALLOWED (Restricted View)", "ALLOWED (Full Access)", "Firebase Auth"],
        ["View Inbound Produce Batches", "ALLOWED (Read-Only)", "ALLOWED (Read-Write)", "Cloud Firestore"],
        ["Approval Status Banner", "SHOWS AMBER WARNING", "HIDDEN / BADGE ACTIVE", "js/wholesaler.js DOM"],
        ["Audit & Certify Farm Stock", "BLOCKED & DISABLED", "FULLY AUTHORIZED", "firestore.rules + UI"],
        ["Receive Transit Handover", "BLOCKED", "FULLY AUTHORIZED", "Order Lifecycle Logic"],
        ["Listed in Admin Approvals Queue", "PENDING QUEUE", "ACCREDITED QUEUE", "admin/approvals.html"]
    ]
    add_styled_table(doc, perm_headers, perm_rows, [1.8, 1.6, 1.6, 1.5])

    # 4. Technical Enforcement & Code Architecture
    add_heading_1(doc, "4. Technical Implementation & Security Code")
    add_body_paragraph(
        doc,
        "The accreditation workflow is enforced at both the database level (Firestore Security Rules) and the client-side presentation layer:"
    )

    add_heading_2(doc, "4.1 Firestore Security Rule Enforcement")
    add_body_paragraph(
        doc,
        "In firestore.rules, stock verification actions require the requester to hold both role == 'wholesaler' AND approved == true. Non-accredited updates are rejected directly by the database engine."
    )

    add_heading_2(doc, "4.2 Client-Side UI Lock (js/wholesaler.js)")
    add_body_paragraph(
        doc,
        "When rendering order tables, js/wholesaler.js evaluates the user's approval status. If approved is false, an amber alert banner is displayed, and the verification action button is disabled with a locked icon and explanatory tooltip: 'Accreditation Pending: Admin approval required to verify stock'."
    )

    # 5. Testing & Verification Scenarios
    add_heading_1(doc, "5. Testing & Verification Demonstration")
    add_body_paragraph(
        doc,
        "The project database is pre-seeded with two authentic Australian wholesalers to allow immediate side-by-side verification:"
    )
    add_bullet_point(doc, "wholesaler@example.com (Liam Wilson - Sydney Central Markets NSW): Pre-seeded as approved: true. Demonstrates an authorized wholesaler actively verifying farm batches.", "1. Approved Wholesaler: ")
    add_bullet_point(doc, "wholesaler.pending@example.com (Matilda Evans - Melbourne Produce Hub VIC): Pre-seeded as approved: false. Demonstrates the pending warning banner and locked verification controls.", "2. Pending Wholesaler: ")
    add_bullet_point(doc, "admin@example.com: Demonstrates navigating to admin/approvals.html, clicking 'Approve Wholesaler' on Matilda, and instantly unlocking her verification powers.", "3. Administrator Approval: ")

    add_screenshot(doc, "19_browser_login_demo_buttons.png", "1-Click Demo Bar on login.html Providing Quick Switching Between Approved and Pending Wholesalers", width_inches=6.0)

    output_path = os.path.join("docx", "wholesaler_approval.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_wholesaler_approval_doc()
