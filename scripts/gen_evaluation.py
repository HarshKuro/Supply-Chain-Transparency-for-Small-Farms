import os
import sys
from docx.shared import Inches, Pt, RGBColor
from doc_utils import (
    init_document, add_title_header, add_heading_1, add_heading_2, add_heading_3,
    add_body_paragraph, add_bullet_point, add_callout,
    add_styled_table, add_screenshot
)

def build_evaluation_doc():
    doc = init_document("Commercial Evaluation & Financial Feasibility Report", "Financial Evaluation")
    add_title_header(
        doc,
        title="Project Commercial Evaluation & Financial Revenue Realization",
        subtitle="Comprehensive financial feasibility report detailing platform charges, real money revenue streams, CAPEX development expenditures, ongoing OPEX, and small farm ROI in Australian Dollars ($ AUD).",
        category_tag="Commercial Feasibility & Financial Model"
    )

    # 1. Executive Summary: Real Money Realization
    add_heading_1(doc, "1. Executive Summary & The Core Financial Question")
    add_body_paragraph(
        doc,
        "A critical question for investors, cooperative leaders, and government grant evaluators is: 'How much real money will this platform generate, what are the platform charges, and how do we sustain commercial profitability while empowering small family farms?'"
    )
    add_body_paragraph(
        doc,
        "AgriTrace operates on a lean, sustainable multi-stream revenue model in Australian Dollars ($ AUD). By eliminating predatory supermarket middlemen (who capture 50–60% markups), AgriTrace charges modest, transparent utility fees that generate reliable recurring operational cash flow while increasing small farmer take-home earnings by +95.5%."
    )

    summary_headers = ["Key Financial Metric", "Value (AUD $)", "Strategic Meaning"]
    summary_rows = [
        ["Initial Development Cost (CAPEX)", "$18,900 AUD", "Lean MVP software build across 220 contractor hours @ standard rates"],
        ["Monthly Cloud & Operating Cost (OPEX)", "$280.00 AUD / month", "Serverless Firebase Blaze, domain, automated backups & triage"],
        ["Monthly Gross Revenue Inflow", "$1,027.00 AUD / month", "Real money collected across 4 diversified platform revenue streams"],
        ["Monthly Net Operating Cash Flow", "$747.00 AUD / month", "Net operating profit after deducting all cloud and operational expenses"],
        ["Annual Net Recurring Profit", "$8,964 AUD / year", "Sustained net cash flow generated at 15-farm pilot capacity"],
        ["Commercial Payback Period", "2.1 Years (or <1 yr w/ grant)", "Full payback on $18,900 AUD CAPEX (11.9 mos with $10k agtech grant)"],
        ["Farmer Income Gain (Take-Home)", "+$10,320 AUD / yr / farm", "+95.5% net margin uplift per small family producer vs supermarket wholesale"]
    ]
    add_styled_table(doc, summary_headers, summary_rows, [2.2, 1.8, 2.5])

    # Clean human-designed visual infographic (NO text signs, NO AI slop)
    add_screenshot(doc, "05_commercial_revenue_model.png", "AgriTrace Commercial Monetization & Real Money Monthly Inflow Infographic ($ AUD)", width_inches=6.2)

    add_screenshot(doc, "14_browser_evaluation_report.png", "AgriTrace Project Financial Evaluation Report Web Interface (evaluation-report.html)", width_inches=6.0)

    # 2. Platform Charges & Monetization Breakdown
    add_heading_1(doc, "2. Platform Charge Streams: How Real Money is Earned")
    add_body_paragraph(
        doc,
        "The platform does not rely on speculative ad clicks or personal data resale. Revenue is earned through four realistic, transparent commercial charges designed for smallholder agriculture:"
    )

    add_heading_2(doc, "Charge Stream 1: Fair-Trade Transaction Commission (3.0%)")
    add_body_paragraph(
        doc,
        "A 3.0% fair-trade processing fee is applied to the gross merchandise value (GMV) of produce ordered through the customer marketplace. Compared to the 50–60% deductions taken by major supermarket chains (Coles and Woolworths), a 3.0% transparent fee is universally embraced by growers."
    )
    add_bullet_point(doc, "Average Regional Order Size: $44.00 AUD (e.g. seasonal vegetable box, 2 punnets strawberries, farm honey).", "Basket Size: ")
    add_bullet_point(doc, "Monthly Transaction Volume: 350 customer orders across 15 participating regional Victorian family farms.", "Volume: ")
    add_bullet_point(doc, "Monthly Gross Merchandise Value (GMV): 350 orders × $44.00 = $15,400.00 AUD / month.", "Gross GMV: ")
    add_bullet_point(doc, "Real Money Revenue: $15,400 × 3.0% = $462.00 AUD / month ($5,544.00 AUD / year).", "Platform Take: ")

    add_heading_2(doc, "Charge Stream 2: Commercial Wholesaler & Buyer Verification Fee ($35.00 AUD / mo)")
    add_body_paragraph(
        doc,
        "Accredited commercial wholesalers, independent grocers, and farm-to-table restaurants pay a modest monthly subscription fee for access to pre-verified batch quality certificates, biosecurity compliance records, and direct bulk purchasing."
    )
    add_bullet_point(doc, "Wholesaler Monthly Fee: $35.00 AUD / month (or $350.00 AUD / year billed annually).", "Pricing: ")
    add_bullet_point(doc, "Pilot Cohort: 8 accredited regional commercial buyers / distributors.", "Client Base: ")
    add_bullet_point(doc, "Real Money Revenue: 8 buyers × $35.00 = $280.00 AUD / month ($3,360.00 AUD / year).", "Platform Take: ")

    add_heading_2(doc, "Charge Stream 3: Farm Cooperative SaaS Tiers")
    add_body_paragraph(
        doc,
        "To protect micro-producers, entry-level farms pay nothing. Commercial family farms and regional packing cooperatives subscribe to digital inventory and traceability tools:"
    )
    add_bullet_point(doc, "Micro Grower Starter Tier (Under 2 hectares): $0.00 / month (Free entry tier to maximize community adoption).", "Tier 1: ")
    add_bullet_point(doc, "Commercial Family Farm Tier: $15.00 AUD / month. 7 subscribing family farms = $105.00 AUD / month.", "Tier 2: ")
    add_bullet_point(doc, "Regional Cooperative / Packing Shed Tier: $45.00 AUD / month. 2 subscribing cooperative hubs = $90.00 AUD / month.", "Tier 3: ")
    add_bullet_point(doc, "Real Money Revenue: $105 + $90 = $195.00 AUD / month ($2,340.00 AUD / year).", "Platform Take: ")

    add_heading_2(doc, "Charge Stream 4: Digital QR Traceability & Batch Verification Stamp ($0.20 AUD / batch)")
    add_body_paragraph(
        doc,
        "A micro-charge of $0.20 AUD is applied per verified batch QR verification seal generated for cold-chain cartons, proving tamper-proof origin and custody."
    )
    add_bullet_point(doc, "Monthly Batch Volume: 450 verified produce batches / crates dispatched.", "Volume: ")
    add_bullet_point(doc, "Real Money Revenue: 450 batches × $0.20 = $90.00 AUD / month ($1,080.00 AUD / year).", "Platform Take: ")

    # Monthly Summary Table
    add_heading_2(doc, "Summary of Monthly Inflow (Real Money)")
    stream_headers = ["Charge Stream", "Pricing Structure", "Monthly Volume / Subscriptions", "Gross Revenue (AUD $)"]
    stream_rows = [
        ["Fair-Trade Transaction Fee", "3.0% of GMV", "$15,400 GMV (350 orders)", "$462.00 AUD"],
        ["Wholesaler Verification Fee", "$35.00 AUD / month", "8 commercial buyers", "$280.00 AUD"],
        ["Farm Cooperative SaaS Tiers", "$15 / $45 AUD / month", "7 farms + 2 cooperatives", "$195.00 AUD"],
        ["Digital QR Batch Stamp Fee", "$0.20 AUD / batch", "450 verified batches", "$90.00 AUD"],
        ["TOTAL MONTHLY REVENUE", "Consolidated Inflow", "All Revenue Channels", "$1,027.00 AUD / mo"]
    ]
    add_styled_table(doc, stream_headers, stream_rows, [2.1, 1.7, 1.6, 1.1])

    # 3. Development Expenditures (CAPEX)
    add_heading_1(doc, "3. Initial Capital Expenditure (CAPEX) — Total: $18,900 AUD")
    add_body_paragraph(
        doc,
        "Initial software development, security hardening, user experience design, and testing were costed based on prevailing Australian engineering contractor rates ($75 – $95 AUD/hr) for a lean, agile MVP:"
    )

    capex_headers = ["Development Phase", "Deliverables & Scope", "Effort", "Cost (AUD $)"]
    capex_rows = [
        ["System Architecture & Schema", "Firestore database collections, role rules, biosecurity flow", "24 hrs", "$2,280 AUD"],
        ["UI/UX Design & Prototyping", "Responsive design system, WCAG 2.1 AA accessibility", "36 hrs", "$3,060 AUD"],
        ["Core Software Engineering", "5 stakeholder portals, real-time Firestore tracking, auth guards", "72 hrs", "$6,480 AUD"],
        ["Quality Assurance & Audit Suite", "31 automated tests (Unit, System, UAT, SEO audit)", "28 hrs", "$2,240 AUD"],
        ["Documentation & User Manuals", "User Manual, Installation Guide, Visual Site Map, Seeder", "20 hrs", "$1,500 AUD"],
        ["Pilot Deployment & Farm Onboarding", "Onboarding 15 pilot farms in Yarra Valley / Mornington", "20 hrs", "$1,600 AUD"],
        ["Contingency & Production Buffer", "10% technical buffer for production hardening", "20 hrs", "$1,740 AUD"],
        ["TOTAL DEVELOPMENT (CAPEX)", "Full Production Turnkey Release", "220 hrs", "$18,900 AUD"]
    ]
    add_styled_table(doc, capex_headers, capex_rows, [1.9, 2.5, 0.9, 1.2])

    # 4. Recurring Operating Costs (OPEX)
    add_heading_1(doc, "4. Monthly Recurring Operating Costs (OPEX) — Total: $280.00 AUD / mo")
    add_body_paragraph(
        doc,
        "AgriTrace operates serverless architecture, drastically minimizing fixed monthly overheads:"
    )

    opex_headers = ["Operating Expense Item", "Vendor / Provider", "Monthly Cost (AUD $)", "Annual Cost (AUD $)"]
    opex_rows = [
        ["Cloud Hosting & Firestore Database", "Google Cloud / Firebase Blaze (Sydney region)", "$42.00 AUD", "$504.00 AUD"],
        ["Domain Name & Cloudflare SSL", "Australian .com.au registrar + Cloudflare DNS", "$6.00 AUD", "$72.00 AUD"],
        ["Transactional SMS & Order Alerts", "Twilio Australia SMS / Email notifications", "$32.00 AUD", "$384.00 AUD"],
        ["Telemetry, Uptime & Daily Backups", "Firestore automated exports & error monitoring", "$20.00 AUD", "$240.00 AUD"],
        ["Part-Time Maintenance & Support", "Part-time developer retainer (2.5 hrs/mo @ $72/hr)", "$180.00 AUD", "$2,160.00 AUD"],
        ["TOTAL MONTHLY OPEX", "All Infrastructure & Operating Services", "$280.00 AUD / mo", "$3,360.00 AUD / yr"]
    ]
    add_styled_table(doc, opex_headers, opex_rows, [2.2, 2.0, 1.2, 1.1])

    # 5. Profit & Loss Statement (P&L) and 3-Year Projections
    add_heading_1(doc, "5. Profit & Loss (P&L) Statement & 3-Year Growth Projections")
    add_body_paragraph(
        doc,
        "Below is the commercial 3-year financial model detailing platform expansion from initial Victorian pilot to state-wide adoption:"
    )

    pnl_headers = ["Financial Item", "Year 1 (Pilot 15 Farms)", "Year 2 (VIC Expansion)", "Year 3 (VIC + NSW Scale)"]
    pnl_rows = [
        ["Participating Family Farms", "15 Farms", "45 Farms", "120 Farms"],
        ["Accredited Wholesalers / Buyers", "8 Buyers", "24 Buyers", "60 Buyers"],
        ["Annual Gross Merchandise Value (GMV)", "$184,800 AUD", "$620,000 AUD", "$1,850,000 AUD"],
        ["Transaction Fees (3.0%)", "$5,544 AUD", "$18,600 AUD", "$55,500 AUD"],
        ["Wholesaler Verification Fees", "$3,360 AUD", "$10,080 AUD", "$25,200 AUD"],
        ["Farm SaaS Subscriptions", "$2,340 AUD", "$9,720 AUD", "$32,400 AUD"],
        ["Digital QR Batch Stamp Fees", "$1,080 AUD", "$4,200 AUD", "$15,300 AUD"],
        ["GROSS ANNUAL REVENUE", "$12,324 AUD", "$42,600 AUD", "$128,400 AUD"],
        ["Less: Annual OPEX Infrastructure", "($3,360 AUD)", "($7,200 AUD)", "($18,000 AUD)"],
        ["NET OPERATING CASH PROFIT", "$8,964 AUD", "$35,400 AUD", "$110,400 AUD"]
    ]
    add_styled_table(doc, pnl_headers, pnl_rows, [2.3, 1.4, 1.4, 1.4])

    # 6. Farmer Economic ROI
    add_heading_1(doc, "6. Small Family Farm Economic Impact & Real Money Gain")
    add_body_paragraph(
        doc,
        "The economic benefit to small family producers is grounded in authentic Australian farm-gate produce economics:"
    )
    add_bullet_point(doc, "Supermarket Duopoly Baseline: A small strawberry grower in the Yarra Valley typically receives $1.80 AUD per 500g punnet from duopoly buyers, while retail consumers are charged $5.50 AUD in store (Grower retains only 32.7% of retail value). Up to 18% of edible stock is cosmetically rejected without payment.", "Duopoly Reality: ")
    add_bullet_point(doc, "AgriTrace Direct Model: The grower lists at $4.20 AUD per punnet (consumers pay $1.30 less than supermarket price). After the 3.0% platform fee ($0.13 AUD) and local transit deduction ($0.55 AUD), the farmer nets $3.52 AUD per punnet — a +95.5% increase in net revenue per unit.", "AgriTrace Model: ")
    add_bullet_point(doc, "Real Money Gain per Farm: For a small family grower harvesting 6,000 units per season, net income rises from $10,800 AUD to $21,120 AUD, delivering +$10,320.00 AUD in extra cash profit directly into the family farm's bank account.", "Annual Farm Gain: ")
    add_bullet_point(doc, "Collective Rural Wealth: Across the 15 pilot farms, over $154,800.00 AUD in collective additional income is retained by local rural producers rather than lost to corporate supermarket margins.", "Cooperative Impact: ")

    add_screenshot(doc, "17_browser_admin_financial_eval.png", "Administrator In-Portal Financial Evaluation Dashboard (admin/evaluation.html)", width_inches=6.0)
    add_screenshot(doc, "16_browser_admin_dashboard.png", "Administrator Platform Dashboard Tracking Cumulative Revenue in Australian Dollars", width_inches=6.0)

    output_path = os.path.join("docx", "evaluation_report.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_evaluation_doc()
