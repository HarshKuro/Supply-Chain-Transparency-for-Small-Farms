import os
import sys
from docx.shared import Inches, Pt, RGBColor
from doc_utils import (
    init_document, add_title_header, add_heading_1, add_heading_2, add_heading_3,
    add_body_paragraph, add_bullet_point, add_callout,
    add_styled_table, add_screenshot
)

def build_admin_integration_doc():
    doc = init_document("Administrator Portal Integration & Platform Governance", "Platform Governance")
    add_title_header(
        doc,
        title="Administrator Portal Integration & Platform Governance",
        subtitle="Architectural specification, macro-oversight dashboards, security boundaries, and multi-role administration for AgriTrace.",
        category_tag="Platform Integration Specification"
    )

    # 1. Executive Context & Integration Purpose
    add_heading_1(doc, "1. Executive Context & The Governance Imperative")
    add_body_paragraph(
        doc,
        "While decentralized agricultural networks provide autonomy to small family growers, complete lack of administrative governance introduces existential operational risks: rogue wholesalers falsely certifying uninspected produce, price collusion, unmonitored cold-chain spoilage disputes, and unverified accounts."
    )
    add_body_paragraph(
        doc,
        "The Administrator Portal (admin/) was fully integrated to provide neutral, macro-level governance over all Australian regional supply nodes, enforcing regulatory compliance, real-time transaction oversight in Australian Dollars ($ AUD), and wholesaler accreditation."
    )

    # 2. Global Navigation & Authentication Routing (REAL IMAGE)
    add_heading_1(doc, "2. Navigation Integration & Authentication Route Architecture")
    add_body_paragraph(
        doc,
        "The Admin Portal has been seamlessly woven into the public information architecture. The diagram below illustrates how authentication routes dispatch users securely into their respective role boundaries:"
    )

    # Embedded high-res visual flowchart (NO text signs)
    add_screenshot(doc, "flow_sitemap_hierarchy.png", "Administrator and Role Authentication Routing Architecture", width_inches=6.2)

    add_bullet_point(doc, "Landing Page Header & Footer: Added direct 'Admin' access links on index.html alongside public documentation links.", "1. Entry Points: ")
    add_bullet_point(doc, "1-Click Demo Toolbar: The login.html portal features a dedicated '🛡️ Admin Portal' quick-login button that signs into admin@example.com with password123.", "2. Rapid Access: ")
    add_bullet_point(doc, "Auth Guarding (js/auth.js): Any attempt by non-admin users (farmers, customers, drivers, wholesalers) to access admin/* triggers immediate interception and rerouting.", "3. Security Fence: ")

    # 3. Admin Operations Dashboard
    add_heading_1(doc, "3. Admin Operations Dashboard (admin/dashboard.html)")
    add_body_paragraph(
        doc,
        "The Administrator Dashboard delivers a high-level operational command center aggregating live platform activity across Australia:"
    )
    add_bullet_point(doc, "Total Platform Revenue: Aggregates real-time order sums in Australian Dollars ($ AUD), displayed prominently as '$XX.XX AUD'.", "Metric 1: ")
    add_bullet_point(doc, "Registered User Breakdown: Counts verified Farmers, Wholesalers, Drivers, and Consumers.", "Metric 2: ")
    add_bullet_point(doc, "Pending Wholesaler Alert Banner: Automatically queries Firestore for wholesalers with approved == false and renders an amber notification badge prompting administrative action.", "Metric 3: ")
    add_bullet_point(doc, "Active Produce Inventory: Tracks total units available across regional farms (e.g. Yarra Valley, Gippsland, Mornington Peninsula).", "Metric 4: ")

    add_callout(
        doc,
        "CRITICAL BUG RESOLUTION: Previously, admin/dashboard.html encountered a JavaScript null reference exception when querying document.getElementById('welcomeName'), interrupting widget initialization. The header template was corrected to include the proper greeting span, ensuring flawless script execution.",
        "Resolved Dashboard Runtime Issue",
        "success"
    )

    add_screenshot(doc, "16_browser_admin_dashboard.png", "Administrator Operations Command Center (admin/dashboard.html) with Real-Time Australian AUD Metrics", width_inches=6.0)

    # 4. Wholesaler Approvals Hub
    add_heading_1(doc, "4. Dedicated Wholesaler Approvals Hub (admin/approvals.html)")
    add_body_paragraph(
        doc,
        "To satisfy Australian food standards and protect small farms, administrators manage wholesaler accreditations through a dedicated interface:"
    )
    add_bullet_point(doc, "Pending Accreditation Queue: Isolates unapproved wholesalers, displaying their Australian Business Name, contact email, and application timestamp.", "Queue A: ")
    add_bullet_point(doc, "Accredited Wholesalers Queue: Lists actively accredited commercial distributors authorized to verify farm shipments.", "Queue B: ")
    add_bullet_point(doc, "1-Click Status Toggling: Administrators can instantly execute toggleWholesalerApproval(userId, true/false), modifying Firestore and logging audit timestamps.", "Action Controls: ")

    add_screenshot(doc, "15_browser_admin_approvals.png", "Wholesaler Accreditation & Approvals Hub (admin/approvals.html)", width_inches=6.0)

    # 5. User Management & Catalog Auditing
    add_heading_1(doc, "5. User Registry & Catalog Governance")
    add_body_paragraph(
        doc,
        "The Admin Portal equips operators with fine-grained inspection capabilities:"
    )
    add_bullet_point(doc, "User Registry (admin/users.html): Complete roster of all platform participants, displaying role badges, phone contacts, and direct approval/revocation buttons for wholesalers.", "Module 1: ")
    add_bullet_point(doc, "Product Catalog Audit (admin/products.html): Comprehensive inspection of all produce batches listed by farmers, verifying fair-trade price boundaries in AUD and stock limits.", "Module 2: ")
    add_bullet_point(doc, "Order Transaction Ledger (admin/orders.html): Global order history providing complete end-to-end transparency across every stage from creation to delivery.", "Module 3: ")
    add_bullet_point(doc, "Financial Feasibility Portal (admin/evaluation.html): Integrated operational view of CAPEX, monthly OPEX, and small farm margin gains.", "Module 4: ")

    add_screenshot(doc, "17_browser_admin_financial_eval.png", "In-Portal Administrative Financial Feasibility & Operating Expense View", width_inches=6.0)

    output_path = os.path.join("docx", "admin_integration.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_admin_integration_doc()
