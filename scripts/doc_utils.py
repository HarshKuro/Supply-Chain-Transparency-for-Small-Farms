import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_HEX = "1B4D20"       # Forest green
SECONDARY_HEX = "2E7D32"     # Leaf green
DARK_SLATE_HEX = "1E293B"    # Slate
LIGHT_BG_HEX = "F8FAF8"      # Tinted light green/gray
BORDER_HEX = "D1D5DB"        # Subtle gray border
ACCENT_AMBER = "B45309"      # Warning amber
ACCENT_BLUE = "1D4ED8"       # Info blue

PRIMARY_RGB = RGBColor(27, 77, 32)
SECONDARY_RGB = RGBColor(46, 125, 50)
DARK_SLATE_RGB = RGBColor(30, 41, 59)
TEXT_MUTED_RGB = RGBColor(100, 116, 139)

def init_document(doc_title, doc_category):
    doc = docx.Document()
    
    # 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        p_head = header.paragraphs[0]
        p_head.text = f"AgriTrace Australia | {doc_category}"
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.runs[0].font.size = Pt(8.5)
        p_head.runs[0].font.color.rgb = TEXT_MUTED_RGB
        p_head.runs[0].font.name = "Calibri"
        
        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.text = f"Supply Chain Transparency for Small Farms • {doc_title} • Confidential / Evaluator Document"
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot.runs[0].font.size = Pt(8.5)
        p_foot.runs[0].font.color.rgb = TEXT_MUTED_RGB
        p_foot.runs[0].font.name = "Calibri"

    return doc

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_title_header(doc, title, subtitle, category_tag):
    # Category tag
    p_cat = doc.add_paragraph()
    p_cat.paragraph_format.space_before = Pt(0)
    p_cat.paragraph_format.space_after = Pt(4)
    run_cat = p_cat.add_run(f"AGRITRACE AUSTRALIA • {category_tag.upper()}")
    run_cat.font.name = "Calibri"
    run_cat.font.size = Pt(9.5)
    run_cat.font.bold = True
    run_cat.font.color.rgb = SECONDARY_RGB

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_RGB

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = DARK_SLATE_RGB

    # Metadata divider box
    meta_table = doc.add_table(rows=1, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_items = [
        ("Platform", "AgriTrace v2.4"),
        ("Jurisdiction", "Australia (AUD $)"),
        ("Date", "October 2026"),
        ("Status", "Production Verified")
    ]
    
    col_widths = [Inches(1.6), Inches(1.6), Inches(1.6), Inches(1.6)]
    for i, (label, val) in enumerate(meta_items):
        cell = meta_table.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        r1 = p.add_run(label.upper() + "\n")
        r1.font.size = Pt(7.5)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_MUTED_RGB
        
        r2 = p.add_run(val)
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = DARK_SLATE_RGB

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = PRIMARY_RGB
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = SECONDARY_RGB
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = DARK_SLATE_RGB
    return h

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_SLATE_RGB
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_SLATE_RGB
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = DARK_SLATE_RGB
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_SLATE_RGB
    return p

def add_callout(doc, text, title="IMPORTANT NOTICE", callout_type="info"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)

    bg_color = "F0FDF4" if callout_type == "success" else ("FEF3C7" if callout_type == "warning" else "F8FAFC")
    title_rgb = PRIMARY_RGB if callout_type == "success" else (RGBColor(180, 83, 9) if callout_type == "warning" else DARK_SLATE_RGB)

    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"■ {title.upper()}\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = title_rgb

    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = DARK_SLATE_RGB

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_diagram_box(doc, title, diagram_lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_background(cell, "0F172A") # Deep dark terminal
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    
    r_title = p.add_run(f"ARCHITECTURE FLOW DIAGRAM: {title}\n")
    r_title.font.name = "Consolas"
    r_title.font.size = Pt(8.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(56, 189, 248) # Cyan

    diagram_text = "\n".join(diagram_lines)
    r_diag = p.add_run(diagram_text)
    r_diag.font.name = "Consolas"
    r_diag.font.size = Pt(8.0)
    r_diag.font.color.rgb = RGBColor(241, 245, 249) # Off-white text

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        if col_widths and col_idx < len(col_widths):
            cell.width = Inches(col_widths[col_idx])
        set_cell_background(cell, PRIMARY_HEX)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header_text)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(rows):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(9.0)
            run.font.color.rgb = DARK_SLATE_RGB

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_screenshot(doc, image_filename, caption, width_inches=6.0):
    img_path = os.path.join("screenshots", image_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figure: {caption}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = TEXT_MUTED_RGB
    else:
        add_callout(doc, f"Screenshot file missing: {image_filename}", "IMAGE NOT FOUND", "warning")
